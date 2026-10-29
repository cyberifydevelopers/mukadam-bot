"""JSON API behind the dashboard's "Stock Reports" tab: every stock report the bot
has made, kept per firm (app/services/stock_report.py writes them to
STOCK_REPORT_DIR/monthly/<firm>/<year-month>/<firm> dd-mm-yy.xlsx + .pdf + a
small .json). The tab lists them by firm and month and opens a report in the
page; the PDF is the sheet exactly as the Excel file looks.

Like the rest of /api this has no login — it is meant for 127.0.0.1 only.
"""

import logging
import re
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse

from app.config import settings
from app.services import stock_report

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/inventory", tags=["inventory"])

_EXTENSIONS = {"xlsx": stock_report.XLSX_MIME, "pdf": stock_report.PDF_MIME}
_TRAILING_DATE = re.compile(r"\s+\d\d-\d\d-\d\d(\s+\(\d+\))?$")


def _root() -> Path:
    return Path(settings.stock_report_dir, "monthly")


def _date_key(date: str) -> str:
    try:
        return datetime.strptime(date, "%d-%m-%Y").strftime("%Y-%m-%d")
    except ValueError:
        return ""


@router.get("")
def list_inventory():
    """The firms with their reports, firms by latest report first, each firm's reports newest first."""
    firms: dict[str, dict] = {}
    for path in sorted(_root().glob("*/*/*.json")):
        state = stock_report._read_state(str(path))
        if not state or not path.with_suffix(".pdf").exists():
            continue
        folder = path.parent.parent.name
        key = state.get("firm") or folder
        customer = state.get("customer") or _TRAILING_DATE.sub("", path.stem)
        firm = firms.setdefault(key, {"key": key, "name": customer, "reports": []})
        firm["reports"].append({
            "ref": f"{folder}/{path.parent.name}/{path.stem}",
            "name": path.stem,
            "date": state.get("date", ""),
            "month": path.parent.name,
            "record": state.get("record"),
            "created": state.get("created", ""),
            "xlsx": path.with_suffix(".xlsx").exists(),
        })
        if state.get("customer"):
            firm["name"] = state["customer"]
    for firm in firms.values():
        firm["reports"].sort(key=lambda r: (_date_key(r["date"]), r["created"]), reverse=True)
        firm["latest"] = firm["reports"][0]["date"]
    ordered = sorted(firms.values(), key=lambda f: (_date_key(f["latest"]), f["name"].lower()), reverse=True)
    return {"firms": ordered}


def _file(ref: str, extension: str) -> Path:
    """The report file named by `ref` (<folder>/<year-month>/<name>), only ever inside the inventory folder."""
    if extension not in _EXTENSIONS:
        raise HTTPException(404, "Unknown file type")
    parts = ref.split("/")
    if len(parts) != 3 or any(p in ("", ".", "..") or "\\" in p for p in parts):
        raise HTTPException(404, "Report not found")
    path = (_root() / parts[0] / parts[1] / f"{parts[2]}.{extension}").resolve()
    if _root().resolve() not in path.parents or not path.is_file():
        raise HTTPException(404, "Report not found")
    return path


@router.get("/file/{ref:path}.{extension}")
def get_file(ref: str, extension: str, download: bool = False):
    """The report as it is: the PDF opens in the browser, the Excel file downloads."""
    path = _file(ref, extension)
    inline = extension == "pdf" and not download
    return FileResponse(
        path, media_type=_EXTENSIONS[extension], filename=path.name,
        content_disposition_type="inline" if inline else "attachment",
    )


@router.delete("/file/{ref:path}")
def delete_report(ref: str):
    """Removes one report (its Excel, PDF and note). The next report of that firm then carries on from the one before it."""
    pdf = _file(ref, "pdf")
    for extension in ("xlsx", "pdf", "json"):
        pdf.with_suffix(f".{extension}").unlink(missing_ok=True)
    logger.info("Stock report %s deleted", ref)
    return {"deleted": ref}


# ---------------------------------------------------------------- one table for a stretch of days

_COLUMNS = [
    ("goods", "Goods"), ("opening", "Opening"), ("received", "Received"), ("delivered", "Delivered"), ("balance", "Balance"),
    ("rate", "Rate"), ("value", "Total value"), ("margin", "Margin %"), ("less_margin", "Less margin"), ("dp", "DP value"),
]
_TEMPLATE_COLUMNS = {
    "opening": "opening", "received": "received", "delivered": "delivered", "balance": "balance", "rate": "rate",
    "value": "value", "margin": "margin", "less_margin": "less_margin", "dp": "dp",
}


def _template(name: str):
    from app.services import stock_templates

    return next((e.template for e in stock_templates.load() if e.template.file == name), None)


def _lines(xlsx: Path, state: dict) -> list[dict]:
    """The stock lines of one finished report, read back from its sheet (figures as LibreOffice calculated them)."""
    import openpyxl

    template = _template(state.get("template", ""))
    if template is None or not xlsx.exists():
        return []
    sheet = openpyxl.load_workbook(xlsx, data_only=True, read_only=True).worksheets[0]
    lines = []
    for r in range(template.first_row, max(template.last_row, state.get("last_row", 0)) + 1):
        goods = next((str(v).strip() for c in (template.description, template.lc_column) if c and (v := sheet[f"{c}{r}"].value) and str(v).strip()), "")
        if not goods:
            continue
        line = {"goods": re.sub(r"\s+", " ", goods)}
        for key, attribute in _TEMPLATE_COLUMNS.items():
            column = getattr(template, attribute, None)
            value = sheet[f"{column}{r}"].value if column else None
            line[key] = value if isinstance(value, (int, float)) else None
        lines.append(line)
    return lines


def combined(firm: str, date_from: str = "", date_to: str = "") -> dict:
    """Every report of a firm between two dates (YYYY-MM-DD, either may be empty), one row per goods line,
    oldest day first."""
    rows, name = [], firm
    for path in sorted(_root().glob("*/*/*.json")):
        state = stock_report._read_state(str(path))
        if not state or (state.get("firm") or path.parent.parent.name) != firm or not path.with_suffix(".xlsx").exists():
            continue
        name = state.get("customer") or name
        day = _date_key(state.get("date", ""))
        if not day or (date_from and day < date_from) or (date_to and day > date_to):
            continue
        for line in _lines(path.with_suffix(".xlsx"), state):
            rows.append({"date": state["date"], "day": day, "created": state.get("created", ""),
                         "ref": f"{path.parent.parent.name}/{path.parent.name}/{path.stem}", "record": state.get("record"), **line})
    rows.sort(key=lambda r: (r["day"], r["created"]))
    used = [(k, label) for k, label in _COLUMNS if k == "goods" or any(r.get(k) is not None for r in rows)]
    return {"firm": firm, "name": name, "from": date_from, "to": date_to, "columns": [{"key": k, "label": l} for k, l in used], "rows": rows}


@router.get("/combined")
def get_combined(firm: str, date_from: str = "", date_to: str = ""):
    return combined(firm, date_from, date_to)


@router.get("/combined.xlsx")
def get_combined_xlsx(firm: str, date_from: str = "", date_to: str = ""):
    """The same table as one Excel file: a row per goods line, date by date."""
    import io

    import openpyxl
    from fastapi.responses import Response
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = combined(firm, date_from, date_to)
    if not data["rows"]:
        raise HTTPException(404, "No reports in these dates")
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Stock report"
    span = f"{data['rows'][0]['date']} to {data['rows'][-1]['date']}"
    sheet["A1"] = f"{data['name']} — stock report, {span}"
    sheet["A1"].font = Font(bold=True, size=13)
    headers = ["Date"] + [c["label"] for c in data["columns"]]
    for i, header in enumerate(headers, 1):
        cell = sheet.cell(row=3, column=i, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(data["rows"], 4):
        sheet.cell(row=r, column=1, value=row["date"])
        for i, column in enumerate(data["columns"], 2):
            value = row.get(column["key"])
            cell = sheet.cell(row=r, column=i, value=value)
            if isinstance(value, (int, float)):
                cell.number_format = "#,##0.00" if column["key"] in ("rate", "value", "less_margin", "dp") else "#,##0.###"
    for i, header in enumerate(headers, 1):
        longest = max([len(str(header))] + [len(str(sheet.cell(row=r, column=i).value or "")) for r in range(4, 4 + len(data["rows"]))])
        sheet.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 44)
    sheet.freeze_panes = "A4"
    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"{_TRAILING_DATE.sub('', data['name']).strip() or 'Stock report'} {data['rows'][0]['date']} to {data['rows'][-1]['date']}.xlsx"
    return Response(
        buffer.getvalue(), media_type=stock_report.XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename.encode("ascii", "ignore").decode() or "stock-report.xlsx"}"'},
    )


# ---------------------------------------------------------------- the bank-approved reports of those days, as one file

def _reports_between(firm: str, date_from: str, date_to: str) -> tuple[str, list[Path]]:
    """(the firm's name, the report files — their .json paths — between two dates, oldest first)."""
    found, name = [], firm
    for path in sorted(_root().glob("*/*/*.json")):
        state = stock_report._read_state(str(path))
        if not state or (state.get("firm") or path.parent.parent.name) != firm:
            continue
        name = state.get("customer") or name
        day = _date_key(state.get("date", ""))
        if day and not (date_from and day < date_from) and not (date_to and day > date_to):
            found.append(((day, state.get("created", "")), path))
    return name, [p for _, p in sorted(found)]


def _span_name(name: str, paths: list[Path]) -> str:
    dates = [stock_report._read_state(str(p)).get("date", "") for p in paths]
    clean = re.sub(r'[<>:"/\|?*]', "", _TRAILING_DATE.sub("", name)).strip() or "Stock report"
    return f"{clean} {dates[0]} to {dates[-1]}"


@router.get("/combined.pdf")
def get_combined_pdf(firm: str, date_from: str = "", date_to: str = "", download: bool = False):
    """Each day's report exactly as the bank approved it (the same pages as the single reports), one after the
    other in date order, in one PDF."""
    import io

    from fastapi.responses import Response
    from pypdf import PdfWriter

    name, paths = _reports_between(firm, date_from, date_to)
    pdfs = [p.with_suffix(".pdf") for p in paths if p.with_suffix(".pdf").exists()]
    if not pdfs:
        raise HTTPException(404, "No reports in these dates")
    writer = PdfWriter()
    for pdf in pdfs:
        writer.append(str(pdf))
    buffer = io.BytesIO()
    writer.write(buffer)
    filename = _span_name(name, paths) + ".pdf"
    disposition = "attachment" if download else "inline"
    return Response(
        buffer.getvalue(), media_type=stock_report.PDF_MIME,
        headers={"Content-Disposition": f'{disposition}; filename="{filename.encode("ascii", "ignore").decode() or "stock-reports.pdf"}"'},
    )


@router.get("/combined.zip")
def get_combined_zip(firm: str, date_from: str = "", date_to: str = ""):
    """The days' Excel files as they are (letterhead and formulas intact), in one zip."""
    import io
    import zipfile

    from fastapi.responses import Response

    name, paths = _reports_between(firm, date_from, date_to)
    sheets = [p.with_suffix(".xlsx") for p in paths if p.with_suffix(".xlsx").exists()]
    if not sheets:
        raise HTTPException(404, "No reports in these dates")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for sheet in sheets:
            archive.write(sheet, sheet.name)
    filename = _span_name(name, paths) + ".zip"
    return Response(
        buffer.getvalue(), media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename.encode("ascii", "ignore").decode() or "stock-reports.zip"}"'},
    )
