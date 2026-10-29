"""The stock report (STOCK REPORT-PLEDGE) sent back on WhatsApp once a
delivery order (or an L/C letter bringing goods in) is confirmed: one of
the company's own report files (app/stock_report_templates/) filled in
completely — each line's description, opening balance, quantity received /
delivered, balance on hand, rate, value, margin and DP value — plus the
date, the bank's city, the godown / customer and the document references.

The sample is filled through LibreOffice (app/services/stock_report_uno.py)
so that everything already in it stays as it is: its TOTAL / drawing-power
formulas, formatting and letterhead. Each line gets formulas too (balance
= opening + received − delivered, value = balance × rate, margin, DP), so
the sheet recalculates when a figure is changed. A rate comes from the
enclosed invoice; a delivery order has none, so its value cells show 0
until one is typed in. LibreOffice also makes the PDF.
"""

import difflib
import json
import logging
import os
import re
import shutil
import subprocess
import threading
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

from app.config import settings
from app.models import InvoiceRecord
from app.services import goods_profiles, stock_templates
from app.services.stock_templates import Template

logger = logging.getLogger(__name__)

PKT = timezone(timedelta(hours=5))  # Pakistan has no DST

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"

_UNO_SCRIPT = Path(__file__).resolve().with_name("stock_report_uno.py")
_office_lock = threading.Lock()  # one LibreOffice at a time: they'd share its profile

_QUANTITY = re.compile(r"^\s*([\d,]*\.?\d+)\s*([A-Za-z.]*)\s*$")
_ZERO_WORDS = {"nil", "-", "—", "none"}
_LC = re.compile(r"\bl\s*/?\s*c\b", re.IGNORECASE)


class StockReportError(Exception):
    pass


class NoFormatError(StockReportError):
    """The firm has no report format (Excel sample) yet: no report is made until one is added."""


@dataclass
class StockReport:
    xlsx_path: str
    pdf_path: str
    customer: str
    template: str
    missing: list[str] = field(default_factory=list)  # what the document didn't give, in words for the approver
    used: list[str] = field(default_factory=list)  # what was filled in from earlier invoices of the same customer's goods


def delivery_order(record: InvoiceRecord) -> dict | None:
    """The delivery-order fields of a confirmed document, or None if it isn't one."""
    order = json.loads(record.raw_extracted_json or "{}").get("delivery_order")
    return order if order and order.get("rows") else None


def _quantity(value: str | None) -> float | None:
    """'10.710 MT' → 10.71; 'NIL' → 0; None if it isn't a plain quantity."""
    value = (value or "").strip()
    if value.lower() in _ZERO_WORDS:
        return 0.0
    m = _QUANTITY.match(value)
    return float(m[1].replace(",", "")) if m else None


_KILOS = {"kg", "kgs", "kilo", "kilos", "kilogram", "kilograms"}


def _unit(value: str | None) -> str:
    """'5600 KGS' → 'kgs'; '' when there is no unit."""
    m = _QUANTITY.match((value or "").strip())
    return m[2].lower().rstrip(".") if m else ""


def _import_conversion(order: dict) -> bool:
    """An invoice in USD is shown in rupees per metric ton (like the approved reports) when a rate is set."""
    return (order.get("currency") or "").strip().upper() == "USD" and settings.usd_pkr_rate > 0


def _customer(order: dict) -> str:
    name = (order.get("customer_name") or "").strip()
    return re.sub(r"^m/s\.?\s*", "", name, flags=re.IGNORECASE) or "Customer"


def _safe_filename(text: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", text).strip(" .")[:80] or "Stock Report"


def bank_of(record: InvoiceRecord, order: dict) -> str:
    """The bank the report is for: the one the client's letter is addressed to, else the one that issued the document."""
    return (order.get("addressee_bank") or record.bank_name or "").strip()


def is_own_format(template: Template, order: dict) -> bool:
    """True if `template` is set up for this firm itself (not borrowed from its bank)."""
    customer = stock_templates.normalise(_customer(order))
    return any(c in customer for c in template.customers)


def format_for(order: dict, bank: str | None = None) -> Template | None:
    """The firm's own format (its uploaded report, matched by name: its rates, margins and goods carry on into
    each new report); failing that, the approved format of its bank — a bank asks for its own layout whichever
    firm the stock is of — else None. With several formats for one bank, the first registered."""
    entries = stock_templates.load()
    customer = stock_templates.normalise(_customer(order))
    for entry in entries:
        if any(c in customer for c in entry.template.customers):
            return entry.template
    if bank:
        for entry in entries:
            if stock_templates.bank_matches(entry.template.banks, bank):
                return entry.template
    return None


def customer_name(record: InvoiceRecord) -> str:
    order = delivery_order(record)
    return _customer(order) if order else ""


def needs_format(record: InvoiceRecord) -> bool:
    """True for a stock document whose firm and bank have no report format yet."""
    order = delivery_order(record)
    return order is not None and format_for(order, bank_of(record, order)) is None


def pick_template(record: InvoiceRecord, order: dict) -> Template:
    bank = bank_of(record, order)
    template = format_for(order, bank)
    if template is None:
        raise NoFormatError(
            f"There is no report format for {_customer(order)}" + (f" or {bank}" if bank else "") + " yet. Add a "
            f"{bank + ' ' if bank else ''}Excel sample (an earlier filled-in report) in the dashboard's Report formats "
            f"tab, then send \"share pdf for {record.id}\"."
        )
    if not is_own_format(template, order):
        logger.info("Stock report #%s: no format for %s, using %s's (%s)", record.id, _customer(order), bank, template.file)
    return template


def _fill_in(text: str, **values: str) -> str:
    """Fills {name} placeholders — and nothing else, so a sample's own braces survive."""
    for key, value in values.items():
        text = text.replace("{" + key + "}", value)
    return text


def _shifted(cell: str, after_row: int, count: int) -> str:
    """A cell below the stock lines, moved down by the rows inserted into them."""
    m = re.match(r"([A-Z]+)(\d+)$", cell)
    return f"{m[1]}{int(m[2]) + count}" if m and int(m[2]) > after_row else cell


def _similarity(a: str, b: str) -> float:
    """How alike two stock descriptions are ('Putty (munds)' and 'Phutty (season 2026-27)'): 1.0 when the
    words of one are all in the other ('Cotton' / 'COTTON BALES'), else the likeness of the spellings."""
    wa, wb = set(a.split()), set(b.split())
    if wa and wb and (wa <= wb or wb <= wa):
        return 1.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def _words(description: str) -> str:
    """Lower-case letters without the bracketed part, and spelled loosely: 'Putti' = 'Putty' = 'Phutty'."""
    text = stock_templates.normalise(re.sub(r"\(.*?\)", " ", description or ""))
    return re.sub(r"(.)\1+", r"\1", text.replace("ph", "p").replace("y", "i"))


def match_lines(rows: list[dict], previous: dict[int, str]) -> list[int | None]:
    """For each line of the document, the row of the previous report holding the same goods (or None)."""
    pairs = sorted(
        (
            (_similarity(_words(item.get("description")), _words(text)), i, row)
            for i, item in enumerate(rows)
            for row, text in previous.items()
        ),
        key=lambda p: -p[0],
    )
    found: list[int | None] = [None] * len(rows)
    taken: set[int] = set()
    for score, i, row in pairs:
        if score >= 0.75 and found[i] is None and row not in taken:
            found[i] = row
            taken.add(row)
    return found


def _extra_columns(template: Template) -> dict[str, str]:
    """Columns the format has for figures the bot doesn't map (drums, weight, weight measurement), found by
    their heading in the row above the stock lines. Only for .xlsx formats; {} when none are found."""
    path = stock_templates.templates_dir() / template.file
    if path.suffix.lower() != ".xlsx" or not path.exists():
        return {}
    found: dict[str, str] = {}
    try:
        sheet = openpyxl.load_workbook(path, read_only=True).worksheets[0]
        header_row = template.first_row - 1
        for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row), ()):
            text = re.sub(r"\s+", " ", str(cell.value or "")).strip().lower()
            column = getattr(cell, "column_letter", None)
            if not text or not column:
                continue
            if "drum" in text:
                found.setdefault("drums", column)
            elif "weight" in text and "measure" in text:
                found.setdefault("weight_measurement", column)
            elif "weight" in text:
                found.setdefault("weight_per_unit", column)
    except Exception:  # noqa: BLE001 — without them the report is still made, just without those cells
        logger.exception("Couldn't read the headings of %s", template.file)
        return {}
    return found


FIGURE_KEYS = {"rate": "rate_per_unit", "margin": "margin_pct", "drums": "drums", "weight_per_unit": "weight_per_unit", "weight_measurement": "weight_measurement"}
FIGURE_NAMES = {"rate": "Rate per unit", "margin": "Margin %", "drums": "No of drums", "weight_per_unit": "Weight per unit", "weight_measurement": "Weight measurement per unit"}


def apply_figures(record: InvoiceRecord, figures: list[tuple[str, str, str]]) -> tuple[list[str], list[str]]:
    """Stores figures the approver typed on WhatsApp — (field, goods name or "", number) — on the document's
    stock lines: the line(s) whose description contains the goods name, or every line when none is named.
    The next report built for this record then carries them. Returns (what was applied, what matched nothing)."""
    raw = json.loads(record.raw_extracted_json or "{}")
    order = raw.get("delivery_order") or {}
    rows = order.get("rows") or []
    applied, unmatched = [], []
    for field, goods, value in figures:
        words = [w for w in re.findall(r"[a-z0-9]+", goods.lower()) if len(w) > 1]
        hit = [row for row in rows if not words or any(w in (row.get("description") or "").lower() for w in words)]
        if not hit and len(rows) == 1:
            hit = rows  # one goods line: the figure can only be for it, however the message worded the goods
        if not hit:
            unmatched.append(f"{FIGURE_NAMES[field]} for \"{goods}\"")
            continue
        for row in hit:
            row[FIGURE_KEYS[field]] = value
        applied.append(f"{FIGURE_NAMES[field]}{f' ({goods})' if goods else ''}: {value}")
    raw["delivery_order"] = order
    record.raw_extracted_json = json.dumps(raw)
    return applied, unmatched


def build_job(template: Template, order: dict, dated: datetime, previous: dict | None = None) -> dict:
    """What stock_report_uno.py writes into the sample. Extra stock lines
    are inserted below the sample's last line, and its TOTAL row's SUM
    ranges are stretched over them, so the totals take them in. Formulas
    the sample already has are kept (see stock_report_uno.py).

    With `previous` ({"last_row", "lines": {row: description}}: the company's last report, which is
    what the job's template then is) the new report carries on from it: every goods line starts
    from its closing balance, the document's own figures go into the lines it names, the rates and
    margins stay, and goods the document doesn't mention stay with no movement. job["state"] is the
    file's state afterwards."""
    rows = order["rows"]
    convert = _import_conversion(order)
    t = template
    job: dict = {"values": [], "merges": [], "text_style": [], "date_cells": [t.date_cell] if t.date_cell else [], "date": f"{dated:%d-%m-%Y}"}
    if previous:
        end = previous["last_row"]
        at = match_lines(rows, previous["lines"])
        free = [r for r in range(t.first_row, end + 1) if r not in previous["lines"]]
        new = [i for i, row in enumerate(at) if row is None]
        added = max(len(new) - len(free), 0)
        for i, row in zip(new, [*free, *range(end + 1, end + added + 1)]):
            at[i] = row
        job["carry"] = {
            "first_row": t.first_row, "last_row": end, "description": t.description, "balance": t.balance,
            "opening": t.opening, "received": t.received, "delivered": t.delivered,
            "formula": "=" + "".join(
                f"{sign}{c}{{r}}" for sign, c in (("+", t.opening), ("+", t.received), ("-", t.delivered)) if c
            ).lstrip("+"),
        }
    else:
        end = t.last_row
        at = [t.first_row + i for i in range(len(rows))]
        added = max(len(rows) - (end - t.first_row + 1), 0)
    if added:
        job.update(insert_after=end, insert_count=added, first_row=t.first_row, last_row=end)
        job["merges"] = [f"{a}{r}:{b}{r}" for r in range(end + 1, end + added + 1) for a, b in t.row_merges]
    job["state"] = {"last_row": end + added}
    mapped = [
        c for c in (
            t.description, t.opening, t.received, t.delivered, t.balance,
            t.rate, t.value, t.dp, t.margin, t.less_margin, t.lc_column,
            t.last_column,
        ) if c
    ]
    # The stock lines and the TOTAL row under them: a figure too wide for its column shrinks to fit instead of showing ###.
    job["fit"] = {
        "first_row": t.first_row, "last_row": end + added + 1,
        "first_column": min(mapped, key=lambda c: (len(c), c)), "last_column": max(mapped, key=lambda c: (len(c), c)),
    }
    if t.clear_sample and not previous:
        job["clear"] = {
            "first_row": t.first_row, "last_row": max(t.last_row, t.first_row + len(rows) - 1),
            "first_column": min(mapped, key=lambda c: (len(c), c)), "last_column": max(mapped, key=lambda c: (len(c), c)),
        }

    if t.margin:  # the "Less Margin @ x%" under the table follows the margins on the lines
        job["margin_label"] = {"column": t.margin, "first_row": t.first_row, "last_row": end + added}
    values = job["values"]
    extra = _extra_columns(t)

    def put(cell: str, value, format_from: str | None = None) -> None:
        values.append([cell, value, format_from] if format_from else [cell, value])

    for i, item in enumerate(rows):
        r = at[i]
        # A goods line the previous report already has keeps its own name, formulas, rate and margin.
        known = bool(previous) and r in previous["lines"]
        description = item.get("description") or ""
        column = t.lc_column if t.lc_column and _LC.search(description) else t.description
        if not known:
            put(f"{column}{r}", description)
            job["text_style"].append([f"{column}{r}", f"{column}{t.first_row}"])  # the sample's own stock-line look

        # Quantities: as stated; a movement not stated is 0, and a balance
        # not stated (goods coming in on an L/C) is worked out in the sheet.
        # A foreign invoice counted in kilos is reported in metric tons (the rate then follows: per ton).
        per = 1.0
        if convert and any(_unit(item.get(k)) in _KILOS for k in ("opening_balance", "quantity_received", "quantity_released", "balance_after")):
            per = 1000.0
        for key, col in (("opening_balance", t.opening), ("quantity_received", t.received), ("quantity_released", t.delivered)):
            if not col:
                continue
            value = (item.get(key) or "").strip()
            number = _quantity(value) if value else 0.0
            if known and (not value or number == 0):
                continue  # the carried opening stays, and a movement not stated stays empty
            put(f"{col}{r}", value if number is None else number / per, f"{col}{t.first_row}")
        if not known:
            balance = (item.get("balance_after") or "").strip()
            if balance and _quantity(balance) is not None:
                put(f"{t.balance}{r}", _quantity(balance) / per, f"{t.balance}{t.first_row}")
            else:
                terms = "".join(
                    f"{sign}{col}{r}" for sign, col in (("+", t.opening), ("+", t.received), ("-", t.delivered)) if col
                ).lstrip("+")
                put(f"{t.balance}{r}", f"={terms}", f"{t.balance}{t.first_row}")

        # Value = balance on hand × rate; then margin and DP value.
        rate = _quantity(item.get("rate_per_unit"))
        if rate is not None:
            if convert:
                rate = round(rate * per * settings.usd_pkr_rate, 2)
            put(f"{t.rate}{r}", rate, f"{t.rate}{t.first_row}")
        # Figures typed in by the approver (drums, weight…) go into their columns, wherever the format has them.
        for key, col in extra.items():
            number = _quantity(str(item.get(key) or ""))
            if number is not None:
                put(f"{col}{r}", number, f"{col}{t.first_row}")
        pct = _quantity(str(item.get("margin_pct") or ""))
        if t.margin and (pct is not None or not known):
            # The margin is the one given for these goods — by the document, the bank, or what was kept for the customer —
            # and never the sample's own: a line without one has an empty margin, and is asked for. Less margin and DP value
            # follow the margin cell, replacing the sample's own formulas ("=L16*5%").
            put(f"{t.margin}{r}", pct if pct is not None else "")
            force = job.setdefault("force", [])
            if t.less_margin:
                put(f"{t.less_margin}{r}", f"={t.value}{r}*{t.margin}{r}/100", f"{t.less_margin}{t.first_row}")
                force.append(f"{t.less_margin}{r}")
            if t.dp:
                put(f"{t.dp}{r}", f"={t.value}{r}-{t.less_margin}{r}" if t.less_margin else f"={t.value}{r}*(100-{t.margin}{r})/100", f"{t.dp}{t.first_row}")
                force.append(f"{t.dp}{r}")
        if known:
            continue
        divisor = f"/{t.value_divisor}" if t.value_divisor != 1 else ""
        put(f"{t.value}{r}", f"={t.balance}{r}*{t.rate}{r}{divisor}", f"{t.value}{t.first_row}")
        if t.dp and not t.margin:  # a format with no margin column: the sample's own margin
            used_pct = t.margin_pct if pct is None else pct
            if used_pct is not None:
                put(f"{t.dp}{r}", f"={t.value}{r}*{100 - used_pct}/100", f"{t.dp}{t.first_row}")

    if previous:
        return job  # the previous report already has its sentence, city and currency labels

    # Only the bank-approved sentence goes in — nothing is appended to it — and it keeps the sample's look.
    location = order.get("pledge_location") or "—"
    put(t.text_cell, _fill_in(t.text, location=location, customer=_customer(order)))
    job["rich"] = {"cell": t.text_cell, "location": location, "customer": _customer(order)}

    branch = (order.get("bank_branch") or "").strip()
    if t.city_cell and "," in branch:
        put(t.city_cell, _fill_in(t.city_text, city=branch.rsplit(",", 1)[1].strip().rstrip(".")))

    currency = (order.get("currency") or "").strip().upper()
    if currency and currency not in ("PKR", "RS", "RS.") and not convert:
        job["replacements"] = [[_shifted(cell, t.last_row, added), r"\bRs\b\.?", currency] for cell in t.currency_cells]
    return job


def _soffice() -> str | None:
    if settings.libreoffice_path:
        return settings.libreoffice_path
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for path in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(path):
            return path
    return None


def _office_python(soffice: str) -> str:
    """LibreOffice's own Python on Windows (it ships with `uno`); elsewhere
    the system python3, with the python3-uno package."""
    bundled = Path(soffice).with_name("python.exe" if os.name == "nt" else "python")
    if bundled.exists():
        return str(bundled)
    # The apt python3-uno package only works with the distro python, not e.g. /usr/local/bin/python3 in a python image.
    return "/usr/bin/python3" if os.path.exists("/usr/bin/python3") else "python3"


def _read_state(path: str) -> dict | None:
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def _company_folders(customer: str, template: Template) -> list[str]:
    """Folders of the company's chain, the one new reports go into first. A customer the format is set up for
    ('Canon Foam' / 'Cannon Foam') shares one chain under the format's name; any other has one under its own
    name, however it is capitalised ('Mian Cotton…' / 'MIAN COTTON…'). The name folder is also searched, for
    reports made before the format folder existed."""
    named = _safe_filename(customer).casefold()
    if any(c in stock_templates.normalise(customer) for c in template.customers):
        return [Path(template.file).stem.casefold(), named]
    return [named]


def _own_report(record_id: int) -> tuple[dict, Path] | None:
    """The newest chain report made for this record: its state and .json path."""
    found = [
        (s, p) for p in Path(settings.stock_report_dir, "monthly").glob("*/*/*.json")
        if (s := _read_state(str(p))) and s.get("record") == record_id and p.with_suffix(".pdf").exists()
    ]
    return max(found, key=lambda sp: sp[0].get("created", "")) if found else None


def _order_key(state: dict) -> tuple[str, str]:
    """Where a report stands in its company's chain: by its date, then by when it was made."""
    try:
        day = datetime.strptime(state["date"], "%d-%m-%Y").strftime("%Y-%m-%d")
    except (KeyError, ValueError):
        day = ""
    return day, state.get("created", "")


def _previous_report(customer: str, template: Template, before: tuple[str, str], record: int | None = None) -> tuple[dict, str] | None:
    """The company's report just before `before` (a date and creation time, see _order_key) — its state
    and .xlsx — if it was made from the same format. A report made from the same document (`record`, read and
    confirmed again) is that document's own earlier try, never what comes before it."""
    states = [(s, p) for name in _company_folders(customer, template)
              for p in Path(settings.stock_report_dir, "monthly", name).glob("*/*.json")
              if (s := _read_state(str(p))) and _order_key(s) < before and (record is None or s.get("record") != record)]
    if not states:
        return None
    state, path = max(states, key=lambda sp: _order_key(sp[0]))
    xlsx = path.with_suffix(".xlsx")
    if state.get("template") != template.file or not xlsx.exists():
        return None  # the company's format changed: a fresh report from the sample
    return state, str(xlsx.resolve())  # LibreOffice needs an absolute path


def _previous_lines(template: Template, state: dict, xlsx: str) -> dict[int, str]:
    """{row: goods description} of the previous report's stock lines."""
    sheet = openpyxl.load_workbook(xlsx, read_only=True).worksheets[0]
    lines = {}
    for r in range(template.first_row, state["last_row"] + 1):
        for column in (template.description, template.lc_column):
            text = sheet[f"{column}{r}"].value if column else None
            if isinstance(text, str) and text.strip():
                lines[r] = text.strip()
                break
    return lines


def render(
    template: Template, order: dict, key: str, dated: datetime | None = None, monthly_record: int | None = None,
    regenerate: bool = False,
) -> StockReport:
    """Fills `template` with the order's lines → an .xlsx and a .pdf under stock_report_dir/<key>/.

    With `monthly_record` (the record's id) the report is one of the company's chain instead
    (stock_report_dir/monthly/<company>/<year-month>/<company> dd-mm-yy.xlsx): a new report starts
    from the company's previous one — its closing balances become the opening ones, and its rates,
    margins and formulas stay — and the new document's figures are put into it. With `regenerate`, a
    report already made for that record is made again in place — from the report that came before it —
    instead of a new one being added to the chain."""
    soffice = _soffice()
    if not soffice:
        raise StockReportError("LibreOffice is not installed on the server (needed to fill the stock report)")

    path = stock_templates.templates_dir() / template.file
    if not path.exists():
        raise StockReportError(f"The report format file {template.file} is missing from {path.parent}")
    customer = _customer(order)
    own = _own_report(monthly_record) if regenerate and monthly_record is not None else None
    if own and not dated:
        dated = datetime.strptime(own[0]["date"], "%d-%m-%Y").replace(tzinfo=PKT)  # no date on the document: it keeps its own
    dated = dated or datetime.now(PKT)
    if monthly_record is None:
        out_dir = os.path.abspath(os.path.join(settings.stock_report_dir, key))
    else:
        out_dir = os.path.abspath(os.path.join(settings.stock_report_dir, "monthly", _company_folders(customer, template)[0], f"{dated:%Y-%m}"))
    os.makedirs(out_dir, exist_ok=True)

    # LibreOffice's Python must not pick up this app's venv.
    env = {k: v for k, v in os.environ.items() if k not in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV")}
    job_path = os.path.join(out_dir, f"job-{os.getpid()}-{threading.get_ident()}.json")
    try:
        with _office_lock:  # also keeps two orders from chaining on the same report at once
            name = f"{_safe_filename(customer)} {dated:%d-%m-%y}"
            base = os.path.join(out_dir, name)
            own_base = str(own[1].resolve().with_suffix("")) if own else None
            if monthly_record is not None:
                n = 1
                while os.path.exists(base + ".xlsx") and base != own_base:  # a second report on the same day is a new file
                    n += 1
                    base = os.path.join(out_dir, f"{name} ({n})")
            previous = _previous_report(
                customer, template, _order_key(own[0]) if own else (f"{dated:%Y-%m-%d}", "￿"), monthly_record
            ) if monthly_record is not None else None
            if (monthly_record is not None and previous is None and template.clear_sample and path.suffix.lower() == ".xlsx"
                    and is_own_format(template, order)):
                # An uploaded sample is a report somebody already made: it is the company's previous report —
                # unless it is another firm's, borrowed for its bank's layout: then it is only wiped and refilled.
                previous = ({"last_row": template.last_row}, str(path))
            lines = _previous_lines(template, *previous) if previous else None

            job = build_job(template, order, dated, {"last_row": previous[0]["last_row"], "lines": lines} if previous and lines else None)
            job.update(
                soffice=soffice,
                # A profile of its own, so it doesn't collide with a LibreOffice window the user has open.
                profile_url=Path(settings.stock_report_dir, ".libreoffice_profile").resolve().as_uri(),
                template=previous[1] if previous and lines else str(path),
                xlsx_out=base + ".xlsx",
                pdf_out=base + ".pdf",
            )
            with open(job_path, "w", encoding="utf-8") as f:
                json.dump(job, f)
            result = subprocess.run(
                [_office_python(soffice), str(_UNO_SCRIPT), job_path], capture_output=True, text=True, timeout=300, env=env
            )
            if result.returncode == 0 and os.path.exists(job["xlsx_out"]) and os.path.exists(job["pdf_out"]):
                if own_base and own_base != base:  # made again under another date: the old files go
                    for extension in (".xlsx", ".pdf", ".json"):
                        if os.path.exists(own_base + extension):
                            os.remove(own_base + extension)
                if monthly_record is not None:
                    with open(base + ".json", "w", encoding="utf-8") as f:
                        json.dump(
                            {"template": template.file, "record": monthly_record, "date": f"{dated:%d-%m-%Y}", "customer": customer,
                             "firm": _company_folders(customer, template)[0],
                             "created": datetime.now(timezone.utc).isoformat(), **job["state"]},
                            f,
                        )
    except (OSError, subprocess.SubprocessError) as exc:
        raise StockReportError(f"LibreOffice couldn't be run: {exc}") from exc
    finally:
        if os.path.exists(job_path):
            os.remove(job_path)
    if result.returncode != 0 or not os.path.exists(job["xlsx_out"]) or not os.path.exists(job["pdf_out"]):
        logger.error("Stock report %s failed: %s", key, (result.stderr or result.stdout)[-2000:])
        raise StockReportError("LibreOffice failed while filling the report (see the server log)")

    logger.info("Stock report %s filled from %s", key, template.file)
    return StockReport(xlsx_path=job["xlsx_out"], pdf_path=job["pdf_out"], customer=customer, template=template.file)


def latest(record: InvoiceRecord) -> StockReport | None:
    """The stock report already made for this record (the newest), or None."""
    own = _own_report(record.id)
    if own:
        state, path = own
        order = delivery_order(record)
        xlsx = path.with_suffix(".xlsx")
        return StockReport(
            xlsx_path=str(xlsx) if xlsx.exists() else "", pdf_path=str(path.with_suffix(".pdf")),
            customer=_customer(order) if order else path.stem, template=state.get("template", ""),
        )
    out_dir = Path(settings.stock_report_dir, str(record.id))
    pdfs = sorted(out_dir.glob("*.pdf"), key=lambda p: p.stat().st_mtime, reverse=True) if out_dir.is_dir() else []
    if not pdfs:
        return None
    xlsx = pdfs[0].with_suffix(".xlsx")
    order = delivery_order(record)
    return StockReport(
        xlsx_path=str(xlsx) if xlsx.exists() else "", pdf_path=str(pdfs[0]),
        customer=_customer(order) if order else pdfs[0].stem, template="",
    )


_DATE_FORMATS = ("%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%y", "%d-%m-%y", "%d/%m/%y", "%d-%b-%y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y", "%d-%B-%Y")


def document_date(record: InvoiceRecord) -> datetime | None:
    """The date the document itself carries ('19.10.2026', '23-Oct-26'); None if it has none or it is a future date."""
    text = (record.txn_date or "").strip()
    for fmt in _DATE_FORMATS:
        try:
            day = datetime.strptime(text, fmt).replace(tzinfo=PKT)
        except ValueError:
            continue
        return day if day <= datetime.now(PKT) else None
    return None


_FIGURE_HEADER = re.compile(r"weight|measure|drum|rate|price|value|margin|dp|drawing", re.IGNORECASE)


def _blank_figure_columns(sheet, template: Template, rows: list[int]) -> list[str]:
    """Headings of the report's figure columns (weight, drums, rate, value, margin, DP…) that have nothing in any
    stock line — those figures are not in the document. Read from the finished sheet, so a rate or margin carried
    on from the previous report doesn't count; zeros (what a formula gives with no rate) count as empty."""
    header_row = template.first_row - 1
    blank = []
    for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row), ()):
        heading = re.sub(r"\s+", " ", str(cell.value or "")).strip()
        if not heading or not _FIGURE_HEADER.search(heading) or not getattr(cell, "column_letter", None):
            continue
        values = (sheet[f"{cell.column_letter}{r}"].value for r in rows)
        if not any(v not in (None, "", 0, 0.0, "0") and str(v).strip() not in ("", "0", "0.0") for v in values):
            blank.append(heading)
    return blank


def missing_data(order: dict, template: Template, report: StockReport, dated: datetime | None) -> list[str]:
    """What the report still lacks because the document didn't give it — read back from the finished sheet, so
    a rate carried on from the previous report doesn't count as missing."""
    notes = []
    try:
        state = _read_state(str(Path(report.xlsx_path).with_suffix(".json"))) or {}
        sheet = openpyxl.load_workbook(report.xlsx_path, data_only=True, read_only=True).worksheets[0]
        rows, no_rate = [], []
        for r in range(template.first_row, max(template.last_row, state.get("last_row", 0)) + 1):
            name = next((str(v).strip() for c in (template.description, template.lc_column) if c and (v := sheet[f"{c}{r}"].value) and str(v).strip()), "")
            if name:
                rows.append(r)
                if not sheet[f"{template.rate}{r}"].value:
                    no_rate.append(name)
        blank = _blank_figure_columns(sheet, template, rows) if rows else []
    except Exception:  # noqa: BLE001 — a note about gaps must never stop the report being sent
        logger.exception("Couldn't check %s for missing data", report.xlsx_path)
        no_rate, blank = [], []
    if blank:
        notes.extend(blank)
    elif no_rate:
        notes.append(f"Rate per unit: {', '.join(no_rate)}. Total Value, Less Margin and Drawing Power could not be worked out")
    if not (order.get("pledge_location") or "").strip():
        notes.append("Godown / pledge location")
    if dated is None:
        notes.append("Date of the document (today's date was used)")
    return notes


def caption(record: InvoiceRecord, report: StockReport) -> str:
    """The text sent with the report: what it is, and anything the document left out."""
    text = f"Stock report for #{record.id} — {report.customer}"
    if report.used:
        lines = "\n".join(f"• {u}" for u in report.used)
        text += f"\n\n♻️ Filled in from what you gave earlier:\n{lines}\n(To change one, reply with e.g. \"margin 6\".)"
    if report.missing:
        lines = "\n".join(f"• {m}" for m in report.missing)
        text += f"\n\n⚠️ These fields are not in the document:\n{lines}\n\nKindly provide them, or fill them in the Excel sheet."
    return text


def build(record: InvoiceRecord, regenerate: bool = False) -> StockReport:
    order = delivery_order(record)
    if order is None:
        raise StockReportError(f"#{record.id} has no delivery-order details to report")
    # A document confirmed again (read again after a correction) replaces its report instead of adding a second one.
    regenerate = regenerate or _own_report(record.id) is not None
    template, dated = pick_template(record, order), document_date(record)
    # What was given once for this customer's goods (margin, weight…) is filled in; what is given now is kept for next time.
    used = goods_profiles.apply(order)
    goods_profiles.remember(order)
    report = render(template, order, str(record.id), dated=dated, monthly_record=record.id, regenerate=regenerate)
    report.missing = missing_data(order, template, report, dated)
    report.used = used
    return report
