"""Turns a saved attachment (PDF / Excel / CSV) into something the LLM can
read: plain text when the file has it, or the decrypted PDF bytes when it's
a scanned PDF (images only) — those go to the model's vision input. An
email with no attachment is saved as a .txt of its headers + body
(gmail_service.save_body_document()) and read here as kind "email".

Password-protected PDFs: which bank sent a PDF is only known after the LLM
has classified the email, which itself benefits from reading the PDF — so
rather than guess the bank first, every configured bank password
(BankProfile.pdf_password_env) is tried until one opens it.
"""

import csv
import glob
import html
import logging
import os
import re
import zipfile
from dataclasses import dataclass, field

from pypdf import PdfReader

from app.bank_registry import BANK_REGISTRY
from app.services.pdf_extractor import NoTextLayerError, PdfExtractionError, extract_text, read_decrypted_bytes

logger = logging.getLogger(__name__)

_MAX_TEXT = 15000  # characters of a spreadsheet/text handed to the LLM
_MAX_ROWS = 400


@dataclass
class Document:
    path: str
    filename: str
    kind: str  # "pdf_text" / "pdf_scanned" / "excel" / "csv" / "word" / "email" / "unreadable"
    text: str | None = None  # for pdf_text / excel / csv / word / email
    password: str | None = None  # the password that opened it, if any
    error: str | None = None  # for "unreadable"
    # For "email" / "word": the pictures in it (logos, photos) — saved files,
    # read by the vision model along with the text.
    image_paths: list[str] = field(default_factory=list)

    @property
    def pdf_bytes(self) -> bytes:
        """Decrypted PDF bytes, for the vision model (scanned PDFs)."""
        return read_decrypted_bytes(self.path, password=self.password)

    def preview(self, chars: int = 1500) -> str:
        """Short description of the content for the email classifier."""
        if self.kind == "pdf_scanned":
            return "[scanned PDF — images only, no text layer]"
        if self.kind == "unreadable":
            return f"[could not be read: {self.error}]"
        return (self.text or "")[:chars]


def _candidate_passwords() -> list[str | None]:
    passwords: list[str | None] = [None]
    for profile in BANK_REGISTRY:
        pw = os.environ.get(profile.pdf_password_env) if profile.pdf_password_env else None
        if pw and pw not in passwords:
            passwords.append(pw)
    return passwords


def _pdf_password(path: str) -> str | None:
    """The password that opens this PDF (None if it isn't encrypted).
    Raises PdfExtractionError when none of the configured ones work."""
    reader = PdfReader(path)
    if not reader.is_encrypted:
        return None
    for pw in _candidate_passwords()[1:]:
        if reader.decrypt(pw):
            return pw
    raise PdfExtractionError(
        "PDF is password-protected and none of the configured bank passwords "
        "(BankProfile.pdf_password_env in app/bank_registry.py) open it"
    )


def _read_pdf(doc: Document) -> None:
    doc.password = _pdf_password(doc.path)
    try:
        doc.text = extract_text(doc.path, password=doc.password)
        doc.kind = "pdf_text"
    except NoTextLayerError:
        doc.kind = "pdf_scanned"


def _rows_to_text(sheet_name: str, rows) -> list[str]:
    lines = [f"## Sheet: {sheet_name}"]
    for i, row in enumerate(rows):
        if i >= _MAX_ROWS:
            lines.append("… (more rows not shown)")
            break
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _read_excel(doc: Document) -> None:
    lines: list[str] = []
    if doc.filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(doc.path)
        for sheet in book.sheets():
            lines += _rows_to_text(sheet.name, (sheet.row_values(r) for r in range(sheet.nrows)))
    else:
        from openpyxl import load_workbook

        book = load_workbook(doc.path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            lines += _rows_to_text(sheet.title, sheet.iter_rows(values_only=True))
        book.close()
    doc.text = "\n".join(lines)[:_MAX_TEXT]
    doc.kind = "excel"


def _read_csv(doc: Document) -> None:
    with open(doc.path, newline="", encoding="utf-8-sig", errors="replace") as f:
        doc.text = "\n".join(_rows_to_text(doc.filename, csv.reader(f)))[:_MAX_TEXT]
    doc.kind = "csv"


def _read_email(doc: Document) -> None:
    with open(doc.path, encoding="utf-8", errors="replace") as f:
        doc.text = f.read()[:_MAX_TEXT]
    message_id = os.path.basename(doc.path).split("_", 1)[0]
    doc.image_paths = sorted(glob.glob(os.path.join(glob.escape(os.path.dirname(doc.path)), f"{glob.escape(message_id)}_img*")))
    doc.kind = "email"


_WORD_IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".webp")
_MAX_WORD_IMAGES = 4


def _word_xml_to_text(xml: str) -> str:
    """WordprocessingML → text: a line per paragraph, a table row's cells
    joined with " | "."""
    xml = re.sub(r"<w:tab/>", "\t", xml)
    # Paragraphs inside a cell stay on the cell's line.
    xml = re.sub(r"<w:tc\b.*?</w:tc>", lambda m: m.group(0).replace("</w:p>", " "), xml, flags=re.S)
    xml = re.sub(r"</w:tc>", " | ", xml)
    xml = re.sub(r"</w:(p|tr)>", "\n", xml)
    text = html.unescape(re.sub(r"<[^>]+>", "", xml))
    lines = [line.strip(" |\t") for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _read_word(doc: Document) -> None:
    """A .docx is a zip of XML: the body, plus headers / footers (where a
    letterhead's bank name usually is) and word/media/ (its logo). The
    pictures are saved next to the file for the vision model."""
    with zipfile.ZipFile(doc.path) as z:
        names = z.namelist()
        parts = [n for n in names if re.fullmatch(r"word/(header\d*|document|footer\d*)\.xml", n)]
        parts.sort(key=lambda n: (0 if "header" in n else 1 if "document" in n else 2, n))
        doc.text = "\n\n".join(_word_xml_to_text(z.read(n).decode("utf-8", errors="replace")) for n in parts)[:_MAX_TEXT]
        media = [n for n in names if n.startswith("word/media/") and n.lower().endswith(_WORD_IMAGE_EXT)]
        for i, name in enumerate(media[:_MAX_WORD_IMAGES], start=1):
            out = f"{doc.path}.media{i}{os.path.splitext(name)[1].lower()}"
            with open(out, "wb") as f:
                f.write(z.read(name))
            doc.image_paths.append(out)
    doc.kind = "word"


def read_document(path: str) -> Document:
    filename = os.path.basename(path).split("_", 1)[-1]  # strip the Gmail message-id prefix
    doc = Document(path=path, filename=filename, kind="unreadable")
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            _read_pdf(doc)
        elif lower.endswith((".xlsx", ".xlsm", ".xls")):
            _read_excel(doc)
        elif lower.endswith(".csv"):
            _read_csv(doc)
        elif lower.endswith(".txt"):
            _read_email(doc)
        elif lower.endswith(".docx"):
            _read_word(doc)
        else:
            doc.error = "unsupported file type"
    except Exception as exc:  # corrupt files, wrong passwords, odd encodings
        doc.kind = "unreadable"
        doc.error = str(exc) or type(exc).__name__
        logger.warning("Could not read %s: %s", filename, doc.error)
    return doc
